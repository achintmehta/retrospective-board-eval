#!/usr/bin/env python3
"""
build_s1_table.py - rebuild Table S1, the per-run supplemental table, as a spreadsheet and,
optionally, as a Word document formatted like the manuscript's other table files.

Usage:   cd analysis && python3 build_s1_table.py            # writes S1_Table.xlsx
         cd analysis && python3 build_s1_table.py --docx     # also writes S1_Table.docx
         cd analysis && python3 build_s1_table.py --docx-only
         (REPO=/path/to/retrospective-board-eval, OUT=/path/S1_Table.xlsx, OUT_DOCX=/path/S1_Table.docx
          override the locations)
Reads:   REPO/<run>/EVALUATION_RUBRIC.md  -> model, harness, effort, tool, per-criterion ratings, score, cost
         ./aesthetic_ratings.csv           -> the holistic 1-5 visual rating per run (editable)
Writes:  S1_Table.xlsx with two sheets, and with --docx an S1_Table.docx with the same two tables:
           "Per-run data"  one row per run: harness, model, effort, UI testing tool, design prompt,
                           score, cost, log10(cost), visual rating
           "Criteria"      one row per run: the 14 criterion ratings (3 = passed first try, 2 = fixed
                           after one corrective prompt, 1 = unresolved), the total, the first-try-perfect
                           flag and the count of criteria rated 2
Needs:   Python 3.8+; openpyxl for the spreadsheet; python-docx for the Word document
         (pip install openpyxl python-docx).

Notes on the labels:
  - log10(cost) is written as a value, not a formula, so it shows in viewers that do not recalculate.
  - The Antigravity harness performed its own interface testing (a browser subagent; in three runs the
    rubric records GPT-OSS 120B as the model it used for that). These are recorded as the rubric states
    them, with the underscore variant normalised, so the column means "UI testing model/tool as
    recorded in the rubric".
  - Effort is left blank where neither the rubric nor the folder name records one (the Antigravity
    Claude and Gemini Flash runs).
"""
import os, re, csv, glob, math, sys, argparse

HERE = os.path.dirname(os.path.abspath(__file__))
def _find_repo():
    """Locate the dataset repo holding <run>/EVALUATION_RUBRIC.md. Works whether this script
    sits inside the repo (analysis/) or beside it (a sibling working folder)."""
    for cand in (os.path.join(HERE, ".."), os.path.join(HERE, "..", "retrospective-board-eval")):
        if glob.glob(os.path.join(cand, "claude_*", "EVALUATION_RUBRIC.md")):
            return os.path.abspath(cand)
    raise SystemExit("dataset not found: set REPO=/path/to/retrospective-board-eval")
REPO = os.environ.get("REPO") or _find_repo()
OUT  = os.environ.get("OUT",  os.path.join(HERE, "S1_Table.xlsx"))
OUT_DOCX = os.environ.get("OUT_DOCX", os.path.join(HERE, "S1_Table.docx"))
RATINGS_CSV = os.path.join(HERE, "aesthetic_ratings.csv")

CAPTION = ("Per-run data for all 90 runs. Part A lists each run's harness, model, reasoning effort, "
           "UI testing model or tool as recorded in its rubric, design prompt condition, functional score "
           "(out of 42), session cost in US dollars as reported by the harness, the base-10 logarithm of that "
           "cost, and the holistic visual rating (1 to 5). The Antigravity harness does not report session "
           "cost. The Qwen cost figures are token-volume estimates at hosted-model rates for locally run "
           "inference, not prices. Part B lists the rating of each of the 14 functional criteria (3 = passed "
           "on the first try, 2 = failed first and was fixed after one corrective prompt, 1 = unresolved), "
           "the total, whether the run was perfect on the first try, and the number of criteria rated 2, "
           "which is the run's corrective-prompt count.")

CRIT_NAMES = ["Local dev environment", "Docker deployment", "Home page", "Board creation", "User identification",
              "Card interaction", "Moving cards", "Commenting", "Realtime: new card", "Realtime: card move",
              "Realtime: new comment", "Data persistence", "Documentation", "CSV export"]
HEAD_A = ["Run folder", "Harness", "Model", "Effort", "UI testing model/tool", "Design prompt",
          "Score (/42)", "Cost (USD)", "log10(cost)", "Visual rating (1-5)"]
HEAD_B = ["Run folder"] + [f"{i}. {n}" for i, n in enumerate(CRIT_NAMES, start=1)] + \
         ["Score (/42)", "First-try perfect", "Criteria rated 2"]
# short criterion labels for the Word table, whose columns are narrow; the full names go in the caption
CRIT_SHORT = ["Local dev", "Docker", "Home", "Board", "User id", "Cards", "Move", "Comment",
              "RT card", "RT move", "RT comment", "Persist", "Docs", "CSV"]
HEAD_B_DOCX = ["Run folder"] + [f"{i}. {n}" for i, n in enumerate(CRIT_SHORT, start=1)] + \
              ["Score (/42)", "First-try perfect", "Criteria rated 2"]

# ------------------------------------------------------------------ data
def field(txt, *labels):
    for lab in labels:
        m = re.search(r'^[ \t>*-]*'+re.escape(lab)+r'[ \t]*:?[ \t]*(.+)$', txt, re.M|re.I)
        if m: return m.group(1).strip().strip('*').strip()
    return ""

def effort_from(folder, em):
    """Effort as recorded in the rubric; otherwise from the folder name for Claude Code runs
    (High unless the name says xHigh or Max). Antigravity runs record no effort setting except
    the two Gemini 3.1 Pro runs, whose folder names carry the agent's high/low thinking level."""
    if em: return em
    f = folder.lower()
    if f.startswith("antigravity"):
        if "_high" in f: return "High"
        if "_low" in f:  return "Low"
        return ""
    if "xhigh" in f: return "xHigh"
    if "max" in f:   return "Max"
    return "High"

def tool_from(raw):
    r = raw.lower().replace("_", " ")
    if "playwright" in r: return "Playwright"
    if "browser" in r:    return "Browser subagent"
    if "gpt-oss" in r:    return "GPT-OSS 120B"
    if "none" in r or r == "": return "None"
    return raw

def load_rows():
    ratings_csv = {}
    with open(RATINGS_CSV, newline="") as f:
        for i, row in enumerate(csv.reader(f)):
            if i == 0 or not row: continue
            ratings_csv[row[0].strip()] = int(row[1])
    rows = []
    for f in sorted(glob.glob(os.path.join(REPO, "*", "EVALUATION_RUBRIC.md"))):
        d = os.path.basename(os.path.dirname(f))
        if d == "template_directory": continue
        txt = open(f, encoding="utf-8", errors="replace").read()
        crits = re.findall(r"\|\s*\*\*(\d+)\*\*\s*\|.*?\|\s*(Pass|Fail)[^|]*\|\s*(\d)\s*\|", txt)
        ratings = {int(c): int(r) for c, _, r in crits}
        if len(ratings) != 14: raise SystemExit(f"{d}: parsed {len(ratings)} criteria, expected 14")
        total = sum(ratings.values())
        declared = re.search(r'Total Score:\**\s*(\d+)\s*/\s*42', txt)
        if declared and int(declared.group(1)) != total:
            raise SystemExit(f"{d}: declared total {declared.group(1)} != sum of ratings {total}")
        costs = re.findall(r'Total cost:\s*\$([0-9][0-9,]*\.[0-9]+)', txt)
        cost = float(costs[-1].replace(",", "")) if costs else None
        rows.append(dict(
            run=d, model=field(txt, "Model Name"),
            harness="Antigravity" if d.startswith("antigravity") else "Claude Code",
            effort=effort_from(d, field(txt, "Effort Mode")),
            tool=tool_from(field(txt, "UI testing model/tool", "UI testing model", "UI testing tool")),
            prompt="Yes" if "antigravity" in d.lower() else ("Abridged" if "abridged" in d.lower() else "No"),
            score=total, cost=cost, log10=(math.log10(cost) if cost else None),
            rating=ratings_csv.get(d), ratings=[ratings[i] for i in range(1, 15)],
            ftp=int(all(v == 3 for v in ratings.values())),
            n2=sum(1 for v in ratings.values() if v == 2),
        ))
    missing = [r["run"] for r in rows if r["rating"] is None]
    if missing: raise SystemExit(f"no visual rating for: {missing}")
    return rows

def row_a(r):  return [r["run"], r["harness"], r["model"], r["effort"], r["tool"], r["prompt"], r["score"], r["cost"], r["log10"], r["rating"]]
def row_b(r):  return [r["run"]] + r["ratings"] + [r["score"], r["ftp"], r["n2"]]

# ------------------------------------------------------------------ xlsx
def write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    hfill = PatternFill("solid", fgColor="305496"); hfont = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    body = Font(name="Arial", size=11); ctr = Alignment(horizontal="center"); left = Alignment(horizontal="left")
    thin = Side(style="thin", color="D9D9D9"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    def header(ws, names):
        ws.append(names)
        for c in range(1, len(names)+1):
            x = ws.cell(1, c); x.fill = hfill; x.font = hfont; x.alignment = ctr; x.border = border
        ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(names))}{len(rows)+1}"
    wb = Workbook(); ws = wb.active; ws.title = "Per-run data"
    header(ws, HEAD_A)
    for i, r in enumerate(rows, start=2):
        for c, v in enumerate(row_a(r), start=1):
            x = ws.cell(i, c, v); x.font = body; x.border = border; x.alignment = ctr if c >= 4 else left
        ws.cell(i, 8).number_format = '"$"#,##0.00'; ws.cell(i, 9).number_format = '0.000'
    for c, wdt in enumerate([50, 12, 20, 9, 22, 14, 11, 12, 12, 18], start=1):
        ws.column_dimensions[get_column_letter(c)].width = wdt
    ws2 = wb.create_sheet("Criteria")
    header(ws2, HEAD_B)
    for i, r in enumerate(rows, start=2):
        for c, v in enumerate(row_b(r), start=1):
            x = ws2.cell(i, c, v); x.font = body; x.border = border; x.alignment = left if c == 1 else ctr
    ws2.column_dimensions["A"].width = 50
    for c in range(2, len(HEAD_B)+1): ws2.column_dimensions[get_column_letter(c)].width = 14
    wb.save(path)
    print("wrote", path)

# ------------------------------------------------------------------ docx
def write_docx(rows, path):
    """Word version in the style of the manuscript's Table_N.docx files: Times New Roman, a bold
    'Table S1.' caption, single-rule borders, a shaded bold header row that repeats on every page,
    US Letter with 2.5 cm margins, landscape so the wide criteria table fits."""
    try:
        from docx import Document
        from docx.shared import Pt, Cm, Twips, Emu
        from docx.enum.section import WD_ORIENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError:
        raise SystemExit("python-docx is required for --docx (pip install python-docx)")

    FONT = "Times New Roman"
    doc = Document()
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None: zoom.set(qn("w:percent"), "100")
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Twips(15840), Twips(12240)          # US Letter, landscape
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.5))
    usable = sec.page_width - sec.left_margin - sec.right_margin

    def run(p, text, bold=False, size=11):
        r = p.add_run(text); r.bold = bold; r.font.name = FONT; r.font.size = Pt(size)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); return r

    def caption(label, text):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
        run(p, label + " ", bold=True); run(p, text)

    def shade(cell, fill="EFEFEF"):
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

    def repeat_header(row):
        trPr = row._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)

    def fmt(v, col):
        if v is None or v == "": return ""
        if col == "cost":  return f"${v:,.2f}"
        if col == "log10": return f"{v:.3f}"
        return str(v)

    def table(head, data, widths, size, left_cols=(0,)):
        t = doc.add_table(rows=1, cols=len(head)); t.style = "Table Grid"
        t.autofit = False            # python-docx writes <w:tblLayout w:type="fixed"/> for this
        # column widths: the relative widths are scaled to the usable page width (an EMU length)
        col_w = [Emu(int(usable * w / sum(widths))) for w in widths]
        for j, w in enumerate(col_w): t.columns[j].width = w
        hdr = t.rows[0]; repeat_header(hdr)
        for j, name in enumerate(head):
            c = hdr.cells[j]; c.width = col_w[j]; shade(c)
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j in left_cols else WD_ALIGN_PARAGRAPH.CENTER
            run(p, name, bold=True, size=size)
        for vals in data:
            cells = t.add_row().cells
            for j, (v, col) in enumerate(vals):
                cells[j].width = col_w[j]
                p = cells[j].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j in left_cols else WD_ALIGN_PARAGRAPH.CENTER
                run(p, fmt(v, col), size=size)
        # tight cell padding, as in the other table files
        for row in t.rows:
            for c in row.cells:
                tcPr = c._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar")
                for side, w in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
                    e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(w)); e.set(qn("w:type"), "dxa"); mar.append(e)
                tcPr.append(mar)
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        return t

    caption("Table S1.", CAPTION)
    caption("Part A.", "Per-run configuration, score, cost and visual rating.")
    cols_a = ["run", "harness", "model", "effort", "tool", "prompt", "score", "cost", "log10", "rating"]
    table(HEAD_A, [list(zip(row_a(r), cols_a)) for r in rows],
          widths=[3300, 900, 1250, 650, 1300, 850, 700, 800, 800, 850], size=8)

    p = doc.add_paragraph(); run(p, "").add_break(WD_BREAK.PAGE)
    caption("Part B.", "Criterion ratings per run (3 = passed first try, 2 = fixed after one corrective prompt, "
            "1 = unresolved), total, first-try-perfect flag (1 = all 14 criteria rated 3) and number of criteria rated 2. "
            "Criteria: " + "; ".join(f"{i} {n}" for i, n in enumerate(CRIT_NAMES, start=1)) + ".")
    cols_b = ["run"] + ["crit"] * 14 + ["score", "ftp", "n2"]
    table(HEAD_B_DOCX, [list(zip(row_b(r), cols_b)) for r in rows],
          widths=[2900] + [590] * 14 + [560, 620, 620], size=7)
    doc.save(path)
    print("wrote", path)

# ------------------------------------------------------------------ main
if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Rebuild Table S1 from the per-run rubric files.")
    ap.add_argument("--docx", action="store_true", help="also write the Word version (S1_Table.docx)")
    ap.add_argument("--docx-only", action="store_true", help="write only the Word version")
    args = ap.parse_args()
    rows = load_rows()
    print("rows:", len(rows))
    if not args.docx_only:
        write_xlsx(rows, OUT)
    if args.docx or args.docx_only:
        write_docx(rows, OUT_DOCX)
